"""게임 카드에 쓰이는 낱말 목록을 엑셀로 뽑습니다.

src/words.js (= tools/build_words.py 가 연구팩에서 만든 파일)를 그대로 읽어
docs/카드_낱말_목록.xlsx 를 만듭니다. 낱말을 바꾸면 build_words.py → 이 스크립트 순서로 다시 돌리세요.

실행: /Users/jini/.hermes/bin/uv run --quiet --with openpyxl python tools/export_words.py
"""
import json
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / 'docs' / '카드_낱말_목록.xlsx'

src = (ROOT / 'src' / 'words.js').read_text(encoding='utf-8')
W = json.loads(src[src.index('{'): src.rindex('}') + 1])
S = W['source']
L1, L2, L3 = W['level1'], W['level2'], W['level3']

# 한 판에 나오는 카드 수 — src/main.js 의 level1Cards / level2Cards / level3Cards 규칙과 같게 셉니다.
l1_cards = len(L1['words'])                                   # 낱말마다 한 장 (10장 중 4장꼴로 여럿)
l1_plural = round(l1_cards * .4)
l2_cards = 10 * 2                                             # 낱말 10개 × 하나/여럿 짝
has_plural_trap = any(t.get('plural') for t in L3['traps'])
both = {w['base'] for w in L3['es']} & {t['w'] for t in L3['traps']}   # 두 묶음에 다 있는 낱말(dress) → 겹치면 한 장
l3_distinct = len({w['base'] for w in L3['es']} | {t['w'] for t in L3['traps']} | {w['base'] for w in L3['regular']})
l3_max = 5 * 2 + 4 + (1 if has_plural_trap else 0) + 6
l3_min = 5 * 2 + 4 + 6 - (1 if both else 0)

HEAD_FILL = PatternFill('solid', fgColor='EEF2F7')
TITLE = Font(bold=True, size=14)
BOLD = Font(bold=True)
MUTED = Font(color='6B7280')
THIN = Side(style='thin', color='D7DCE3')
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical='top', wrap_text=True)
CENTER = Alignment(horizontal='center', vertical='top')


def table(ws, top, headers, rows, widths, center_cols=()):
    """머리줄 + 데이터 줄을 쓰고 테두리·필터·틀 고정까지."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, GRID
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r, row in enumerate(rows, top + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = GRID
            cell.alignment = CENTER if c in center_cols else WRAP
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + c)].width = w
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    ws.auto_filter.ref = f'A{top}:{chr(64 + len(headers))}{top + len(rows)}'


wb = Workbook()

# ── 요약
ws = wb.active
ws.title = '요약'
ws['A1'] = 'Phonics Halli Galli — 카드 낱말 목록'
ws['A1'].font = TITLE
ws['A2'] = f"출처: {S['pack']}"
ws['A3'] = f"기준: {S['sheet']}"
ws['A4'] = f"3단계 함정 낱말: {S['traps']}"
ws['A5'] = f"3단계 그냥 -s: {S['level3Regular']}"
for r in range(2, 6):
    ws[f'A{r}'].font = MUTED
summary = [
    ['1단계', '5세', S['books']['5'], len(L1['words']),
     f'{l1_cards}장 — 낱말마다 한 장, 그중 {l1_plural}장은 여럿(-s)', len(L1['excluded'])],
    ['2단계', '6세', S['books']['6'], len(L2['words']),
     f'{l2_cards}장 — 낱말 10개를 골라 하나/여럿 짝으로', len(L2['excluded'])],
    ['3단계', '7세', S['books']['7'],
     f"{l3_distinct} (-es 짝 {len(L3['es'])} · 함정 {len(L3['traps'])} · 그냥 -s {len(L3['regular'])}"
     + (f" — {', '.join(sorted(both))}는 -es 짝과 함정에 모두 있어 한 번만 셈)" if both else ')'),
     f'{l3_min}~{l3_max}장 — -es 짝 5개 + 함정 4개 + 그냥 -s 6장', len(L3['excluded'])],
]
table(ws, 7, ['단계', '나이', '교재', '쓰는 낱말', '한 판에 나오는 카드', '뺀 낱말'], summary,
      [9, 7, 16, 30, 46, 9], center_cols=(1, 2, 6))
ws.freeze_panes = None
ws.auto_filter.ref = None
ws['A12'] = '낱말은 판마다 섞어서 뽑아요. 1단계는 모든 낱말이 매 판 나오고, 2·3단계는 위 개수만큼 골라요.'
ws['A12'].font = MUTED


# ── 1·2단계: 아이 화면에 나오는 한글 뜻(하나/여럿)까지
def level_rows(words):
    return [[i, w['base'], w['plural'], f"{w['ko']} 한 {w['unit']}", f"{w['ko']} 여러 {w['unit']}", w['n'], w['src']]
            for i, w in enumerate(words, 1)]


for name, L in (('1단계 (5세)', L1), ('2단계 (6세)', L2)):
    ws = wb.create_sheet(name)
    table(ws, 1, ['번호', '하나 카드', '여럿 카드', '그림 뜻 (하나)', '그림 뜻 (여럿)', '교재 출현 수', '나온 쪽'],
          level_rows(L['words']), [6, 12, 12, 18, 18, 11, 60], center_cols=(1, 6))

# ── 3단계: -es 짝 · 함정 · 그냥 -s
rows = []
for w in L3['es']:
    rows.append(['-es 짝', w['base'], w['plural'], w['n'], w['src'], ''])
for t in L3['traps']:
    note = f"예문: {t['sentence']}"
    if t.get('plural'):
        note += f"  /  {t['plural']}: {t['pluralSrc']}"
    rows.append(['함정 (s로 끝나는 하나)', t['w'], t.get('plural') or '— (하나 카드만)', t['n'], t['src'], note])
for w in L3['regular']:
    rows.append(['그냥 -s', w['base'], w['plural'], w['n'], w['src'], ''])
rows = [[i, *r] for i, r in enumerate(rows, 1)]
ws = wb.create_sheet('3단계 (7세)')
table(ws, 1, ['번호', '구분', '하나 카드', '여럿 카드', '교재 출현 수', '나온 쪽', '비고'],
      rows, [6, 22, 12, 16, 11, 52, 60], center_cols=(1, 5))

# ── 뺀 낱말과 이유
ws = wb.create_sheet('뺀 낱말')
out = [[lv, x['w'], x['why']] for lv, L in (('1단계 (5세)', L1), ('2단계 (6세)', L2), ('3단계 (7세)', L3)) for x in L['excluded']]
table(ws, 1, ['단계', '낱말', '뺀 이유'], out, [13, 12, 70], center_cols=(1,))

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(OUT)
print(f"1단계 {len(L1['words'])} · 2단계 {len(L2['words'])} · 3단계 {l3_distinct} ({len(L3['es'])}+{len(L3['traps'])}+{len(L3['regular'])}, 겹침 {sorted(both)}) · 뺀 낱말 {len(out)}")
print(f'한 판 카드: 1단계 {l1_cards} · 2단계 {l2_cards} · 3단계 {l3_min}~{l3_max}')
